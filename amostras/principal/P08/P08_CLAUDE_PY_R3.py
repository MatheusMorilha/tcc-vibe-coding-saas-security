"""
Exportação dos dados cadastrais de clientes de uma empresa (tenant) para XLSX.

Pontos de segurança e robustez:
- O tenant (empresa_id) vem SEMPRE do usuário autenticado, nunca de parâmetro
  da requisição, evitando acesso a dados de outras empresas (IDOR).
- Exige permissão explícita de exportação (dados pessoais - LGPD).
- Células de texto são gravadas como string literal, nunca como fórmula
  (proteção contra formula injection: =, +, -, @).
- Geração em modo "write-only" com leitura em lotes: memória constante,
  mesmo com muitos clientes.
- Limite máximo de linhas por exportação.
- Toda exportação é registrada em log de auditoria.
"""
from __future__ import annotations

import logging
from datetime import datetime
from tempfile import TemporaryFile
from typing import IO, Iterator
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

# --- Ajuste estes imports à estrutura do seu projeto -------------------------
from app.core.deps import get_current_user, get_db
from app.models.cliente import Cliente
from app.models.usuario import Usuario
# -----------------------------------------------------------------------------

logger = logging.getLogger("audit.clientes_export")

router = APIRouter(prefix="/api/v1/clientes", tags=["Clientes"])

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
PERMISSAO_EXPORTAR = "clientes:exportar"
MAX_LINHAS_EXPORTACAO = 100_000
TAMANHO_LOTE = 1_000
LIMITE_CARACTERES_CELULA = 32_767  # limite do Excel
FUSO_EXPORTACAO = ZoneInfo("America/Sao_Paulo")

# (título da coluna, atributo do modelo Cliente, largura da coluna)
COLUNAS: list[tuple[str, str, int]] = [
    ("ID", "id", 10),
    ("Nome / Razão social", "nome", 36),
    ("Nome fantasia", "nome_fantasia", 30),
    ("CPF / CNPJ", "documento", 20),
    ("E-mail", "email", 34),
    ("Telefone", "telefone", 18),
    ("Logradouro", "logradouro", 34),
    ("Número", "numero", 10),
    ("Complemento", "complemento", 20),
    ("Bairro", "bairro", 22),
    ("Cidade", "cidade", 22),
    ("UF", "uf", 6),
    ("CEP", "cep", 12),
    ("Ativo", "ativo", 8),
    ("Cadastrado em", "criado_em", 20),
]


# ----------------------------------------------------------------------------
# Autorização
# ----------------------------------------------------------------------------
def exigir_permissao_exportacao(
    usuario: Usuario = Depends(get_current_user),
) -> Usuario:
    """Garante que o usuário autenticado pode exportar dados de clientes."""
    if not usuario.tem_permissao(PERMISSAO_EXPORTAR):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Você não tem permissão para exportar clientes.",
        )
    # Ponto de extensão: validar aqui se o plano da empresa inclui exportação.
    return usuario


# ----------------------------------------------------------------------------
# Geração da planilha
# ----------------------------------------------------------------------------
def _celula(ws, valor):
    """Converte um valor Python em célula segura para o XLSX."""
    if valor is None:
        return None

    if isinstance(valor, bool):
        return "Sim" if valor else "Não"

    if isinstance(valor, datetime):
        if valor.tzinfo is not None:  # Excel não suporta datetime com fuso
            valor = valor.astimezone(FUSO_EXPORTACAO).replace(tzinfo=None)
        cel = WriteOnlyCell(ws, value=valor)
        cel.number_format = "DD/MM/YYYY HH:MM"
        return cel

    if isinstance(valor, str):
        texto = ILLEGAL_CHARACTERS_RE.sub("", valor)[:LIMITE_CARACTERES_CELULA]
        cel = WriteOnlyCell(ws, value=texto)
        # Força tipo texto: impede que "=...", "+...", "-...", "@..." sejam
        # interpretados como fórmula pelo Excel/LibreOffice.
        cel.data_type = "s"
        return cel

    return valor  # int, float, Decimal etc.


def _gerar_xlsx(db: Session, filtro: tuple, destino: IO[bytes]) -> None:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Clientes")

    # Em modo write-only, formatações de coluna devem vir antes das linhas.
    for indice, (_, _, largura) in enumerate(COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(indice)].width = largura
    ws.freeze_panes = "A2"

    fonte = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="305496")
    alinhamento = Alignment(horizontal="center", vertical="center")

    cabecalho = []
    for titulo, _, _ in COLUNAS:
        cel = WriteOnlyCell(ws, value=titulo)
        cel.font, cel.fill, cel.alignment = fonte, fundo, alinhamento
        cabecalho.append(cel)
    ws.append(cabecalho)

    # Seleciona apenas as colunas necessárias e lê em lotes (cursor streaming).
    stmt = (
        select(*[getattr(Cliente, attr) for _, attr, _ in COLUNAS])
        .where(*filtro)
        .order_by(Cliente.id)
        .execution_options(yield_per=TAMANHO_LOTE)
    )
    for linha in db.execute(stmt):
        ws.append([_celula(ws, valor) for valor in linha])

    wb.save(destino)
    wb.close()


def _iterar_arquivo(arquivo: IO[bytes], tamanho_bloco: int = 64 * 1024) -> Iterator[bytes]:
    """Envia o arquivo em blocos e garante a remoção do temporário ao final."""
    try:
        while bloco := arquivo.read(tamanho_bloco):
            yield bloco
    finally:
        arquivo.close()  # TemporaryFile é apagado ao fechar


# ----------------------------------------------------------------------------
# Rota
# ----------------------------------------------------------------------------
@router.get(
    "/exportar",
    summary="Exporta os clientes da empresa para XLSX",
    response_class=StreamingResponse,
    responses={
        200: {"content": {XLSX_MEDIA_TYPE: {}}, "description": "Arquivo XLSX"},
        403: {"description": "Sem permissão de exportação"},
        413: {"description": "Volume de dados acima do limite por exportação"},
    },
)
def exportar_clientes(
    db: Session = Depends(get_db),
    usuario: Usuario = Depends(exigir_permissao_exportacao),
) -> StreamingResponse:
    # Isolamento multilocatário: o escopo vem do usuário autenticado.
    empresa_id = usuario.empresa_id
    filtro = (
        Cliente.empresa_id == empresa_id,
        Cliente.excluido_em.is_(None),  # ignora registros com exclusão lógica
    )

    total = db.scalar(select(func.count()).select_from(Cliente).where(*filtro)) or 0
    if total > MAX_LINHAS_EXPORTACAO:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=(
                f"A exportação excede o limite de {MAX_LINHAS_EXPORTACAO} "
                "registros. Entre em contato com o suporte."
            ),
        )

    arquivo = TemporaryFile()
    try:
        _gerar_xlsx(db, filtro, arquivo)
        tamanho = arquivo.tell()
        arquivo.seek(0)
    except Exception:
        arquivo.close()
        logger.exception(
            "Falha ao gerar exportação de clientes",
            extra={"empresa_id": str(empresa_id), "usuario_id": str(usuario.id)},
        )
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Não foi possível gerar a exportação.",
        )

    logger.info(
        "Exportação de clientes realizada",
        extra={
            "empresa_id": str(empresa_id),
            "usuario_id": str(usuario.id),
            "total_registros": total,
        },
    )

    carimbo = datetime.now(FUSO_EXPORTACAO).strftime("%Y%m%d_%H%M%S")
    nome_arquivo = f"clientes_{carimbo}.xlsx"

    return StreamingResponse(
        _iterar_arquivo(arquivo),
        media_type=XLSX_MEDIA_TYPE,
        headers={
            "Content-Disposition": f'attachment; filename="{nome_arquivo}"',
            "Content-Length": str(tamanho),
            "Cache-Control": "no-store",  # dados pessoais: não cachear
            "X-Content-Type-Options": "nosniff",
        },
    )