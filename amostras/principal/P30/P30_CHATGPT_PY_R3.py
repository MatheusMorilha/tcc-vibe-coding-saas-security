from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.auth import get_current_user
from app.database import get_db
from app.models import Client, Company, User


router = APIRouter(prefix="/api/v1/clients", tags=["clients"])


@router.post("/import")
def import_clients(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    """
    Importa clientes da planilha XLSX da empresa do usuário autenticado.

    Colunas esperadas:
        name
        document
        email
        phone
    """

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Apenas arquivos XLSX são aceitos.",
        )

    # O tenant é obtido exclusivamente do contexto autenticado.
    # Nunca deve ser recebido do arquivo ou de um parâmetro enviado pelo cliente.
    company = db.scalar(
        select(Company).where(Company.id == current_user.company_id)
    )

    if company is None:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Empresa não encontrada.",
        )

    try:
        contents = file.file.read()
        workbook = load_workbook(
            filename=BytesIO(contents),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Não foi possível ler a planilha.",
        ) from exc

    worksheet = workbook.active

    rows = worksheet.iter_rows(values_only=True)

    try:
        headers = next(rows)
    except StopIteration:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="A planilha está vazia.",
        )

    normalized_headers = [
        str(value).strip().lower() if value is not None else ""
        for value in headers
    ]

    required_columns = {"name", "document", "email", "phone"}

    if not required_columns.issubset(set(normalized_headers)):
        missing = sorted(required_columns - set(normalized_headers))
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={
                "message": "Colunas obrigatórias ausentes.",
                "missing_columns": missing,
            },
        )

    column_index = {
        column: normalized_headers.index(column)
        for column in required_columns
    }

    processed = 0
    imported = 0
    skipped = 0
    errors: list[dict[str, Any]] = []

    for row_number, row in enumerate(rows, start=2):
        processed += 1

        def get_value(column: str) -> str:
            index = column_index[column]
            value = row[index] if index < len(row) else None
            return str(value).strip() if value is not None else ""

        name = get_value("name")
        document = get_value("document")
        email = get_value("email")
        phone = get_value("phone")

        # Ignora linhas completamente vazias.
        if not any((name, document, email, phone)):
            skipped += 1
            continue

        validation_errors = []

        if not name:
            validation_errors.append("name é obrigatório")

        if not document:
            validation_errors.append("document é obrigatório")

        if not email:
            validation_errors.append("email é obrigatório")

        if validation_errors:
            errors.append(
                {
                    "row": row_number,
                    "errors": validation_errors,
                }
            )
            continue

        # A busca de duplicidade contém obrigatoriamente o company_id.
        # Assim, um cliente de outra empresa nunca será considerado
        # pertencente ao tenant atual.
        existing_client = db.scalar(
            select(Client).where(
                Client.company_id == company.id,
                Client.document == document,
            )
        )

        if existing_client is not None:
            skipped += 1
            errors.append(
                {
                    "row": row_number,
                    "errors": ["Cliente com este documento já existe."],
                }
            )
            continue

        client = Client(
            company_id=company.id,
            name=name,
            document=document,
            email=email,
            phone=phone,
        )

        try:
            # Savepoint por registro: uma linha inválida não desfaz
            # os registros válidos já processados.
            with db.begin_nested():
                db.add(client)
                db.flush()

            imported += 1

        except IntegrityError:
            # Também protege contra condições de corrida ou constraints
            # únicas existentes no banco.
            skipped += 1
            errors.append(
                {
                    "row": row_number,
                    "errors": ["Não foi possível inserir o cliente; registro duplicado ou inválido."],
                }
            )

    db.commit()
    workbook.close()

    return {
        "status": "completed",
        "summary": {
            "processed": processed,
            "imported": imported,
            "skipped": skipped,
            "errors": len(errors),
        },
        "errors": errors,
    }
