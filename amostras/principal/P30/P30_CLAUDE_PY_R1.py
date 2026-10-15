"""
POST /api/v1/clientes/importar
Importação em lote de planilha (CSV ou XLSX) de clientes, isolada por empresa (tenant).

Dependências: fastapi, sqlalchemy>=2.0, psycopg[binary], openpyxl, python-multipart
"""
import csv
import io
import logging
import re
from dataclasses import dataclass, field
from typing import Iterator

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.orm import Session

# Ajuste para os módulos reais do projeto
from app.auth import UsuarioAutenticado, exige_permissao   # usuário vem do token, nunca do payload
from app.db import get_db
from app.models import Assinatura, Cliente, LogAuditoria

logger = logging.getLogger("importacao_clientes")
router = APIRouter(prefix="/api/v1/clientes", tags=["clientes"])

# --------------------------------------------------------------------------- #
# Limites e constantes
# --------------------------------------------------------------------------- #
MAX_BYTES = 5 * 1024 * 1024        # 5 MB por arquivo
MAX_LINHAS_ABSOLUTO = 20_000       # teto global, independente do plano
TAMANHO_LOTE = 500                 # linhas por transação
MAX_ERROS_DETALHADOS = 100         # evita resposta gigante
COLUNAS_OBRIGATORIAS = {"nome", "email"}
COLUNAS_ACEITAS = {"nome", "email", "documento", "telefone"}

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
TIPOS_PERMITIDOS = {
    "text/csv",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


# --------------------------------------------------------------------------- #
# Schemas de resposta
# --------------------------------------------------------------------------- #
class ErroLinha(BaseModel):
    linha: int
    motivo: str


class ResumoImportacao(BaseModel):
    total_lidas: int
    importadas: int
    duplicadas_ignoradas: int
    invalidas: int
    bloqueadas_por_limite_do_plano: int
    erros: list[ErroLinha]
    erros_truncados: bool


@dataclass
class Contadores:
    lidas: int = 0
    importadas: int = 0
    duplicadas: int = 0
    invalidas: int = 0
    bloqueadas: int = 0
    erros: list[ErroLinha] = field(default_factory=list)

    def erro(self, linha: int, motivo: str) -> None:
        self.invalidas += 1
        if len(self.erros) < MAX_ERROS_DETALHADOS:
            self.erros.append(ErroLinha(linha=linha, motivo=motivo))


# --------------------------------------------------------------------------- #
# Leitura do arquivo (streaming, sem carregar tudo em dicts)
# --------------------------------------------------------------------------- #
async def ler_arquivo_limitado(arquivo: UploadFile) -> bytes:
    """Lê o upload em pedaços, abortando se passar do limite (não confia em Content-Length)."""
    buf = bytearray()
    while chunk := await arquivo.read(64 * 1024):
        buf.extend(chunk)
        if len(buf) > MAX_BYTES:
            raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                                f"Arquivo excede {MAX_BYTES // (1024 * 1024)} MB.")
    if not buf:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Arquivo vazio.")
    return bytes(buf)


def _normaliza_cabecalho(h) -> str:
    return str(h or "").strip().lower()


def iterar_linhas(nome: str, conteudo: bytes) -> Iterator[tuple[int, dict]]:
    """Gera (numero_da_linha_na_planilha, {coluna: valor}). Cabeçalho é a linha 1."""
    nome = (nome or "").lower()
    if nome.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        ws = wb.active
        linhas = ws.iter_rows(values_only=True)
        cab = [_normaliza_cabecalho(c) for c in next(linhas, [])]
        for n, row in enumerate(linhas, start=2):
            if row is None or all(c in (None, "") for c in row):
                continue
            yield n, dict(zip(cab, row))
        wb.close()
    elif nome.endswith(".csv"):
        try:
            texto = conteudo.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = conteudo.decode("latin-1")
        amostra = texto[:2048]
        try:
            dialeto = csv.Sniffer().sniff(amostra, delimiters=",;\t")
        except csv.Error:
            dialeto = csv.excel
        leitor = csv.reader(io.StringIO(texto), dialeto)
        cab = [_normaliza_cabecalho(c) for c in next(leitor, [])]
        for n, row in enumerate(leitor, start=2):
            if not any(c.strip() for c in row):
                continue
            yield n, dict(zip(cab, row))
    else:
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
                            "Formato não suportado. Envie .csv ou .xlsx.")


# --------------------------------------------------------------------------- #
# Validação / sanitização de uma linha
# --------------------------------------------------------------------------- #
def _texto(v) -> str:
    return "" if v is None else str(v).strip()


def _neutraliza_formula(v: str) -> str:
    """Mitiga CSV/Formula Injection caso o dado seja reexportado para planilha."""
    return "'" + v if v and v[0] in ("=", "+", "-", "@", "\t", "\r") else v


def validar_linha(bruta: dict) -> tuple[dict | None, str | None]:
    nome = _texto(bruta.get("nome"))
    email = _texto(bruta.get("email")).lower()
    documento = re.sub(r"\D", "", _texto(bruta.get("documento"))) or None
    telefone = _texto(bruta.get("telefone")) or None

    if not nome or len(nome) > 200:
        return None, "Nome obrigatório (máx. 200 caracteres)."
    if not EMAIL_RE.match(email) or len(email) > 254:
        return None, "E-mail inválido."
    if documento and len(documento) not in (11, 14):
        return None, "Documento deve ter 11 (CPF) ou 14 (CNPJ) dígitos."
    if telefone and len(telefone) > 30:
        return None, "Telefone muito longo."

    return {
        "nome": _neutraliza_formula(nome),
        "email": email,
        "documento": documento,
        "telefone": telefone,
    }, None


# --------------------------------------------------------------------------- #
# Regras do plano
# --------------------------------------------------------------------------- #
def obter_capacidade_restante(db: Session, tenant_id: int) -> int:
    """Quantos clientes ainda cabem no plano ativo do tenant."""
    assinatura = db.execute(
        select(Assinatura).where(
            Assinatura.tenant_id == tenant_id,
            Assinatura.status == "ativa",
        )
    ).scalar_one_or_none()

    if assinatura is None:
        raise HTTPException(status.HTTP_402_PAYMENT_REQUIRED,
                            "Nenhuma assinatura ativa para esta empresa.")

    atuais = db.execute(
        select(func.count()).select_from(Cliente).where(Cliente.tenant_id == tenant_id)
    ).scalar_one()

    if assinatura.limite_clientes is None:      # plano ilimitado
        return MAX_LINHAS_ABSOLUTO
    return max(assinatura.limite_clientes - atuais, 0)


# --------------------------------------------------------------------------- #
# Rota
# --------------------------------------------------------------------------- #
@router.post("/importar", response_model=ResumoImportacao)
async def importar_clientes(
    arquivo: UploadFile = File(...),
    usuario: UsuarioAutenticado = Depends(exige_permissao("clientes:importar")),
    db: Session = Depends(get_db),
):
    # O tenant SEMPRE vem da identidade autenticada; nunca de parâmetro, header ou planilha.
    tenant_id = usuario.tenant_id

    if arquivo.content_type not in TIPOS_PERMITIDOS and not (arquivo.filename or "").lower().endswith((".csv", ".xlsx")):
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "Tipo de arquivo não permitido.")

    conteudo = await ler_arquivo_limitado(arquivo)
    capacidade = obter_capacidade_restante(db, tenant_id)
    c = Contadores()

    try:
        linhas = iterar_linhas(arquivo.filename, conteudo)
        primeira = next(linhas, None)
        if primeira is None:
            raise HTTPException(status.HTTP_400_BAD_REQUEST, "Planilha sem dados.")

        cabecalho = set(primeira[1].keys())
        faltando = COLUNAS_OBRIGATORIAS - cabecalho
        if faltando:
            raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                                f"Colunas obrigatórias ausentes: {', '.join(sorted(faltando))}.")

        def todas():
            yield primeira
            yield from linhas

        vistos_no_arquivo: set[str] = set()
        lote: list[tuple[int, dict]] = []

        def processar_lote() -> None:
            nonlocal capacidade
            if not lote:
                return
            emails = [d["email"] for _, d in lote]

            # Duplicidade contra a base, ESCOPADA no tenant
            existentes = set(db.execute(
                select(Cliente.email).where(
                    Cliente.tenant_id == tenant_id,
                    Cliente.email.in_(emails),
                )
            ).scalars())

            novos = []
            for n, d in lote:
                if d["email"] in existentes:
                    c.duplicadas += 1
                elif capacidade <= 0:
                    c.bloqueadas += 1
                else:
                    novos.append({**d, "tenant_id": tenant_id, "criado_por": usuario.id})
                    capacidade -= 1

            if novos:
                # ON CONFLICT protege contra corrida (índice único em tenant_id + email)
                stmt = (
                    pg_insert(Cliente)
                    .values(novos)
                    .on_conflict_do_nothing(index_elements=["tenant_id", "email"])
                    .returning(Cliente.id)
                )
                inseridos = len(db.execute(stmt).all())
                c.importadas += inseridos
                c.duplicadas += len(novos) - inseridos
                capacidade += len(novos) - inseridos   # devolve vaga das que conflitaram
            db.commit()                                 # um commit por lote
            lote.clear()

        for n, bruta in todas():
            c.lidas += 1
            if c.lidas > MAX_LINHAS_ABSOLUTO:
                raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                                    f"Máximo de {MAX_LINHAS_ABSOLUTO} linhas por importação.")

            dados, erro = validar_linha(bruta)
            if erro:
                c.erro(n, erro)
                continue
            if dados["email"] in vistos_no_arquivo:
                c.duplicadas += 1
                continue
            vistos_no_arquivo.add(dados["email"])

            lote.append((n, dados))
            if len(lote) >= TAMANHO_LOTE:
                processar_lote()
        processar_lote()

    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        logger.exception("Falha na importação | tenant=%s usuario=%s", tenant_id, usuario.id)
        # Mensagem genérica: não vaza detalhes internos nem de outros tenants
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR,
                            "Erro ao processar a planilha. Lotes já gravados foram preservados.")

    # Auditoria
    db.add(LogAuditoria(
        tenant_id=tenant_id, usuario_id=usuario.id, acao="clientes.importar",
        detalhes={"arquivo": arquivo.filename, "lidas": c.lidas, "importadas": c.importadas,
                  "duplicadas": c.duplicadas, "invalidas": c.invalidas, "bloqueadas": c.bloqueadas},
    ))
    db.commit()

    return ResumoImportacao(
        total_lidas=c.lidas,
        importadas=c.importadas,
        duplicadas_ignoradas=c.duplicadas,
        invalidas=c.invalidas,
        bloqueadas_por_limite_do_plano=c.bloqueadas,
        erros=c.erros,
        erros_truncados=c.invalidas > len(c.erros),
    )