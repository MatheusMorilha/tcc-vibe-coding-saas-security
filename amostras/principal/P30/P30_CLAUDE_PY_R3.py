"""
POST /customers/import

Importa em lote uma planilha (CSV ou XLSX) de clientes para a empresa (tenant)
do usuário autenticado e devolve um resumo do processamento.

Premissas de isolamento multilocatário:
  * O tenant_id vem EXCLUSIVAMENTE do usuário autenticado (token/sessão).
    Qualquer coluna "tenant_id"/"empresa_id" na planilha é ignorada.
  * Toda consulta e toda escrita é filtrada/carimbada com esse tenant_id.
  * Duplicidade é avaliada apenas dentro do tenant (um mesmo e-mail pode
    existir em empresas diferentes).
  * Recomenda-se na base: UNIQUE (tenant_id, email) e UNIQUE (tenant_id, document).
"""
from __future__ import annotations

import csv
import io
import logging
import re
import unicodedata
from typing import Iterator

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

# Ajuste aos módulos reais do projeto
from app.auth import get_current_user          # retorna usuário com .tenant_id, .id, .permissions
from app.db import get_db                      # fornece a Session
from app.models import Customer, Subscription  # Subscription: tenant_id, status, max_customers

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/customers", tags=["customers"])

# ----------------------------------------------------------------------------
# Limites de segurança / desempenho
# ----------------------------------------------------------------------------
MAX_FILE_BYTES = 5 * 1024 * 1024        # 5 MB
MAX_ROWS_HARD = 20_000                  # teto absoluto por importação
BATCH_SIZE = 500                        # registros por lote/savepoint
MAX_ERRORS_REPORTED = 200               # evita resposta gigante
ALLOWED_EXT = (".csv", ".xlsx")

# Cabeçalhos aceitos (normalizados) -> campo interno. Colunas desconhecidas são ignoradas.
HEADER_ALIASES = {
    "nome": "name", "name": "name", "razao_social": "name", "cliente": "name",
    "email": "email", "e_mail": "email",
    "cpf": "document", "cnpj": "document", "cpf_cnpj": "document", "documento": "document",
    "document": "document",
    "telefone": "phone", "celular": "phone", "phone": "phone",
    "observacoes": "notes", "observacao": "notes", "notes": "notes",
}

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


# ----------------------------------------------------------------------------
# Schemas de resposta
# ----------------------------------------------------------------------------
class ErroLinha(BaseModel):
    linha: int
    motivo: str


class ResumoImportacao(BaseModel):
    total_linhas: int
    importados: int
    duplicados: int
    invalidos: int
    rejeitados_por_limite_do_plano: int
    erros: list[ErroLinha]
    erros_truncados: bool


# ----------------------------------------------------------------------------
# Utilitários de normalização / validação
# ----------------------------------------------------------------------------
def _normalizar_cabecalho(valor: object) -> str:
    s = unicodedata.normalize("NFKD", str(valor or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


def _texto(valor: object, max_len: int) -> str:
    """Limpa texto livre: remove controles e neutraliza CSV/Formula Injection."""
    if valor is None:
        return ""
    s = "".join(ch for ch in str(valor) if ch.isprintable()).strip()[:max_len]
    if s and s[0] in "=+-@":
        s = "'" + s
    return s


def _so_digitos(valor: object) -> str:
    return re.sub(r"\D", "", str(valor or ""))


def _documento_valido(doc: str) -> bool:
    """Valida CPF (11) ou CNPJ (14) pelos dígitos verificadores."""
    if len(doc) not in (11, 14) or len(set(doc)) == 1:
        return False
    if len(doc) == 11:
        for n in (9, 10):
            soma = sum(int(doc[i]) * (n + 1 - i) for i in range(n))
            if (soma * 10 % 11) % 10 != int(doc[n]):
                return False
        return True
    for n, pesos in ((12, [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]),
                     (13, [6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])):
        soma = sum(int(doc[i]) * pesos[i] for i in range(n))
        dv = 0 if soma % 11 < 2 else 11 - soma % 11
        if dv != int(doc[n]):
            return False
    return True


def _validar_linha(bruto: dict) -> tuple[dict | None, str | None]:
    nome = _texto(bruto.get("name"), 150)
    email = str(bruto.get("email") or "").strip().lower()
    doc = _so_digitos(bruto.get("document"))
    fone = _so_digitos(bruto.get("phone"))

    if not nome:
        return None, "nome é obrigatório"
    if not email and not doc:
        return None, "informe ao menos e-mail ou documento"
    if email and (len(email) > 254 or not EMAIL_RE.match(email)):
        return None, "e-mail inválido"
    if doc and not _documento_valido(doc):
        return None, "CPF/CNPJ inválido"
    if fone and not 10 <= len(fone) <= 13:
        return None, "telefone inválido"

    return {
        "name": nome,
        "email": email or None,
        "document": doc or None,
        "phone": fone or None,
        "notes": _texto(bruto.get("notes"), 500) or None,
    }, None


# ----------------------------------------------------------------------------
# Leitura da planilha (streaming; devolve (nº da linha, dict por campo interno))
# ----------------------------------------------------------------------------
def _mapear(cabecalhos: list[object]) -> dict[int, str]:
    mapa = {}
    for i, h in enumerate(cabecalhos):
        campo = HEADER_ALIASES.get(_normalizar_cabecalho(h))
        if campo and campo not in mapa.values():
            mapa[i] = campo
    return mapa


def _ler_csv(conteudo: bytes) -> Iterator[tuple[int, dict]]:
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = conteudo.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    leitor = csv.reader(io.StringIO(texto), dialect)
    mapa: dict[int, str] = {}
    for n, linha in enumerate(leitor, start=1):
        if n == 1:
            mapa = _mapear(linha)
            continue
        if any(c.strip() for c in linha):
            yield n, {campo: linha[i] for i, campo in mapa.items() if i < len(linha)}


def _ler_xlsx(conteudo: bytes) -> Iterator[tuple[int, dict]]:
    from openpyxl import load_workbook  # import tardio: dependência só deste formato

    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)  # data_only: não executa fórmulas
    try:
        mapa: dict[int, str] = {}
        for n, linha in enumerate(wb.active.iter_rows(values_only=True), start=1):
            if n == 1:
                mapa = _mapear(list(linha))
                continue
            if any(c not in (None, "") for c in linha):
                yield n, {campo: linha[i] for i, campo in mapa.items() if i < len(linha)}
    finally:
        wb.close()


# ----------------------------------------------------------------------------
# Gravação de um lote
# ----------------------------------------------------------------------------
def _gravar_lote(
    db: Session,
    tenant_id: int,
    lote: list[tuple[int, dict]],
    vagas: int | None,
    registrar_erro,
) -> tuple[int, int, int]:
    """Retorna (importados, duplicados, rejeitados_por_limite)."""
    emails = [r["email"] for _, r in lote if r["email"]]
    docs = [r["document"] for _, r in lote if r["document"]]

    # Duplicidade contra a base, SEMPRE restrita ao tenant
    existentes_email, existentes_doc = set(), set()
    if emails:
        existentes_email = set(db.scalars(
            select(Customer.email).where(Customer.tenant_id == tenant_id, Customer.email.in_(emails))))
    if docs:
        existentes_doc = set(db.scalars(
            select(Customer.document).where(Customer.tenant_id == tenant_id, Customer.document.in_(docs))))

    novos, duplicados = [], 0
    for linha, reg in lote:
        if (reg["email"] in existentes_email) or (reg["document"] in existentes_doc):
            duplicados += 1
            registrar_erro(linha, "cliente já cadastrado nesta empresa")
        else:
            novos.append((linha, reg))

    # Limite do plano
    rejeitados = 0
    if vagas is not None and len(novos) > vagas:
        for linha, _ in novos[max(vagas, 0):]:
            registrar_erro(linha, "limite de clientes do plano atingido")
        rejeitados = len(novos) - max(vagas, 0)
        novos = novos[: max(vagas, 0)]

    if not novos:
        return 0, duplicados, rejeitados

    # Inserção em lote dentro de um savepoint
    try:
        with db.begin_nested():
            db.bulk_insert_mappings(Customer, [{**reg, "tenant_id": tenant_id} for _, reg in novos])
        return len(novos), duplicados, rejeitados
    except IntegrityError:
        # Corrida com outra requisição/violação de unicidade: reprocessa linha a linha
        importados = 0
        for linha, reg in novos:
            try:
                with db.begin_nested():
                    db.add(Customer(**reg, tenant_id=tenant_id))
                    db.flush()
                importados += 1
            except IntegrityError:
                duplicados += 1
                registrar_erro(linha, "cliente já cadastrado nesta empresa")
        return importados, duplicados, rejeitados


# ----------------------------------------------------------------------------
# Rota
# ----------------------------------------------------------------------------
@router.post("/import", response_model=ResumoImportacao, status_code=status.HTTP_200_OK)
def importar_clientes(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    # 1) Autorização
    if "customers:import" not in getattr(user, "permissions", []):
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Sem permissão para importar clientes.")

    tenant_id = user.tenant_id  # única fonte de verdade do tenant

    # 2) Validação do arquivo (extensão + tamanho, lendo em partes)
    nome = (file.filename or "").lower()
    if not nome.endswith(ALLOWED_EXT):
        raise HTTPException(status.HTTP_415_UNSUPPORTED_MEDIA_TYPE, "Envie um arquivo .csv ou .xlsx.")
    conteudo = bytearray()
    while chunk := file.file.read(64 * 1024):
        conteudo += chunk
        if len(conteudo) > MAX_FILE_BYTES:
            raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Arquivo excede 5 MB.")
    if not conteudo:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Arquivo vazio.")

    # 3) Assinatura ativa + trava do tenant (serializa importações concorrentes,
    #    evitando estourar o limite do plano por corrida)
    assinatura = db.scalar(
        select(Subscription).where(Subscription.tenant_id == tenant_id).with_for_update()
    )
    if assinatura is None or assinatura.status != "active":
        raise HTTPException(status.HTTP_402_PAYMENT_REQUIRED, "Assinatura inativa.")

    vagas: int | None = None
    if assinatura.max_customers is not None:
        atuais = db.scalar(select(func.count()).select_from(Customer).where(Customer.tenant_id == tenant_id))
        vagas = max(assinatura.max_customers - atuais, 0)

    # 4) Processamento
    leitor = _ler_csv(bytes(conteudo)) if nome.endswith(".csv") else _ler_xlsx(bytes(conteudo))

    erros: list[ErroLinha] = []
    truncou = False

    def registrar_erro(linha: int, motivo: str) -> None:
        nonlocal truncou
        if len(erros) < MAX_ERRORS_REPORTED:
            erros.append(ErroLinha(linha=linha, motivo=motivo))
        else:
            truncou = True

    total = importados = duplicados = invalidos = rejeitados = 0
    vistos_email: set[str] = set()
    vistos_doc: set[str] = set()
    lote: list[tuple[int, dict]] = []

    def descarregar() -> None:
        nonlocal importados, duplicados, rejeitados, vagas, lote
        if not lote:
            return
        ok, dup, rej = _gravar_lote(db, tenant_id, lote, vagas, registrar_erro)
        importados += ok
        duplicados += dup
        rejeitados += rej
        if vagas is not None:
            vagas -= ok
        lote = []

    try:
        for numero, bruto in leitor:
            total += 1
            if total > MAX_ROWS_HARD:
                raise HTTPException(
                    status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                    f"Máximo de {MAX_ROWS_HARD} linhas por importação.",
                )

            reg, motivo = _validar_linha(bruto)
            if reg is None:
                invalidos += 1
                registrar_erro(numero, motivo)
                continue

            # Duplicidade dentro do próprio arquivo
            if (reg["email"] and reg["email"] in vistos_email) or (
                reg["document"] and reg["document"] in vistos_doc
            ):
                duplicados += 1
                registrar_erro(numero, "duplicado dentro da planilha")
                continue
            if reg["email"]:
                vistos_email.add(reg["email"])
            if reg["document"]:
                vistos_doc.add(reg["document"])

            lote.append((numero, reg))
            if len(lote) >= BATCH_SIZE:
                descarregar()

        descarregar()
        db.commit()  # única confirmação: libera a trava do tenant

    except HTTPException:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        # Não logar conteúdo da planilha (dados pessoais); apenas contexto técnico
        logger.exception("Falha na importação de clientes tenant=%s user=%s", tenant_id, user.id)
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            "Não foi possível processar a planilha. Verifique o formato e tente novamente.",
        )

    logger.info(
        "Importação concluída tenant=%s user=%s total=%d importados=%d",
        tenant_id, user.id, total, importados,
    )

    # 5) Resumo
    return ResumoImportacao(
        total_linhas=total,
        importados=importados,
        duplicados=duplicados,
        invalidos=invalidos,
        rejeitados_por_limite_do_plano=rejeitados,
        erros=erros,
        erros_truncados=truncou,
    )