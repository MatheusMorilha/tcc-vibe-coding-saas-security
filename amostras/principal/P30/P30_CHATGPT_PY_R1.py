
from io import BytesIO
from typing import Generator
import re
import unicodedata

from fastapi import (
    FastAPI,
    Depends,
    File,
    HTTPException,
    UploadFile,
    status,
)
from openpyxl import load_workbook
from sqlalchemy import (
    String,
    Integer,
    ForeignKey,
    UniqueConstraint,
    create_engine,
    select,
)
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import (
    DeclarativeBase,
    Mapped,
    Session,
    mapped_column,
    sessionmaker,
)


# --------------------------------------------------
# Configuração do banco de dados
# --------------------------------------------------

DATABASE_URL = "sqlite:///./saas.db"

engine = create_engine(
    DATABASE_URL,
    connect_args={"check_same_thread": False},
)

SessionLocal = sessionmaker(
    bind=engine,
    autoflush=False,
    expire_on_commit=False,
)


class Base(DeclarativeBase):
    pass


# --------------------------------------------------
# Modelos
# --------------------------------------------------

class User(Base):
    __tablename__ = "users"

    id: Mapped[int] = mapped_column(primary_key=True)
    tenant_id: Mapped[int] = mapped_column(
        Integer, nullable=False, index=True
    )
    can_import_customers: Mapped[bool] = mapped_column(
        default=False, nullable=False
    )


class Customer(Base):
    __tablename__ = "customers"
    __table_args__ = (
        UniqueConstraint(
            "tenant_id",
            "email",
            name="uq_customer_tenant_email",
        ),
    )

    id: Mapped[int] = mapped_column(primary_key=True)
    tenant_id: Mapped[int] = mapped_column(
        ForeignKey("tenants.id"), nullable=False, index=True
    )
    name: Mapped[str] = mapped_column(String(150), nullable=False)
    email: Mapped[str] = mapped_column(String(254), nullable=False)
    phone: Mapped[str | None] = mapped_column(
        String(30), nullable=True
    )


# Em uma aplicação real, Tenant e User já devem existir
# no modelo de dados da plataforma.
class Tenant(Base):
    __tablename__ = "tenants"

    id: Mapped[int] = mapped_column(primary_key=True)


Base.metadata.create_all(bind=engine)

app = FastAPI(title="API de Importação de Clientes")


# --------------------------------------------------
# Dependências
# --------------------------------------------------

def get_db() -> Generator[Session, None, None]:
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def get_current_user(
    db: Session = Depends(get_db),
) -> User:
    """
    Substitua pela dependência de autenticação JWT
    já utilizada pela plataforma.

    Ela deve validar o token e retornar o usuário
    autenticado consultado no banco.
    """
    raise HTTPException(
        status_code=status.HTTP_401_UNAUTHORIZED,
        detail="Configure a autenticação JWT da aplicação.",
    )


def require_import_permission(
    user: User = Depends(get_current_user),
) -> User:
    if not user.can_import_customers:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Usuário sem permissão para importar clientes.",
        )
    return user


# --------------------------------------------------
# Validação dos dados
# --------------------------------------------------

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_ROWS = 10_000

REQUIRED_COLUMNS = {"nome", "email", "telefone"}


def normalize_header(value: object) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(
        char for char in text
        if not unicodedata.combining(char)
    )
    return text


def validate_row(row: dict) -> dict:
    name = str(row.get("nome") or "").strip()
    email = str(row.get("email") or "").strip().lower()
    phone = str(row.get("telefone") or "").strip() or None

    if not name or len(name) > 150:
        raise ValueError("Nome obrigatório ou maior que 150 caracteres.")

    if len(email) > 254 or not re.fullmatch(
        r"[^@\s]+@[^@\s]+\.[^@\s]+", email
    ):
        raise ValueError("E-mail inválido.")

    if phone and len(phone) > 30:
        raise ValueError("Telefone maior que 30 caracteres.")

    return {
        "name": name,
        "email": email,
        "phone": phone,
    }


# --------------------------------------------------
# Rota de importação
# --------------------------------------------------

@app.post("/api/v1/customers/import")
def import_customers(
    file: UploadFile = File(...),
    user: User = Depends(require_import_permission),
    db: Session = Depends(get_db),
):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=400,
            detail="Envie uma planilha Excel com extensão .xlsx.",
        )

    # A leitura é limitada para evitar uploads excessivos.
    content = file.file.read(MAX_FILE_SIZE + 1)

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="O arquivo excede o limite de 10 MB.",
        )

    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
        worksheet = workbook.active

        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            raise HTTPException(
                status_code=400,
                detail="A planilha está vazia.",
            )

        normalized_headers = [
            normalize_header(value) for value in headers
        ]

        # Aceita telefone opcional, mas exige nome e e-mail.
        if not {"nome", "email"}.issubset(set(normalized_headers)):
            raise HTTPException(
                status_code=400,
                detail="A planilha deve conter as colunas nome e email.",
            )

        header_indexes = {
            header: index
            for index, header in enumerate(normalized_headers)
            if header
        }

        total = 0
        imported = 0
        skipped = 0
        rejected = 0
        errors = []

        try:
            # Uma transação externa para a importação.
            with db.begin():
                for row_number, values in enumerate(rows, start=2):
                    # Ignora linhas inteiramente vazias.
                    if not any(
                        value is not None and str(value).strip()
                        for value in values
                    ):
                        continue

                    total += 1

                    if total > MAX_ROWS:
                        raise HTTPException(
                            status_code=413,
                            detail=f"Limite de {MAX_ROWS} registros excedido.",
                        )

                    raw_data = {
                        column: (
                            values[index]
                            if index < len(values)
                            else None
                        )
                        for column, index in header_indexes.items()
                        if column in REQUIRED_COLUMNS
                        or column == "nome"
                    }

                    try:
                        customer_data = validate_row(raw_data)

                        # O tenant_id vem exclusivamente do usuário
                        # autenticado, nunca da planilha.
                        with db.begin_nested():
                            existing = db.scalar(
                                select(Customer.id).where(
                                    Customer.tenant_id == user.tenant_id,
                                    Customer.email
                                    == customer_data["email"],
                                )
                            )

                            if existing is not None:
                                skipped += 1
                                continue

                            customer = Customer(
                                tenant_id=user.tenant_id,
                                **customer_data,
                            )
                            db.add(customer)
                            db.flush()

                        imported += 1

                    except ValueError as exc:
                        rejected += 1
                        errors.append({
                            "linha": row_number,
                            "motivo": str(exc),
                        })

                    except IntegrityError:
                        # Trata duplicidade concorrente protegida
                        # pela restrição UNIQUE do banco.
                        skipped += 1

            return {
                "mensagem": "Importação concluída.",
                "resumo": {
                    "total_processado": total,
                    "importados": imported,
                    "ignorados": skipped,
                    "rejeitados": rejected,
                },
                "erros": errors,
            }

        except HTTPException:
            raise
        except SQLAlchemyError:
            raise HTTPException(
                status_code=500,
                detail="Falha no banco de dados. A transação foi revertida.",
            )

    finally:
        if "workbook" in locals():
            workbook.close()
        file.file.close()
